#!/usr/bin/env python3
"""Export all datasets from the configured Provena instance to an Excel audit spreadsheet."""

from __future__ import annotations

import argparse
import asyncio
import os
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import urlparse

_project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_project_root))

import boto3
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from provenaclient import ProvenaClient, Config
from provenaclient.auth.implementations import OfflineFlow
from provenaclient.auth.manager import Log
from ProvenaInterfaces.DataStoreAPI import CredentialsRequest
from ProvenaInterfaces.RegistryAPI import (
    FilterOptions,
    GeneralListRequest,
    NoFilterSubtypeListRequest,
    QueryRecordTypes,
    SortOptions,
    SortType,
)
from ProvenaInterfaces.RegistryModels import ItemSubType

from server import provena_runtime as pr

HANDLE_URL_PREFIX = "https://hdl.handle.net/"
COLUMNS = [
    "id",
    "displayname",
    "description",
    "format",
    "date created",
    "date modified",
    "licence",
    "dataset size",
    "dataset owner",
    "dataset custodian",
    "version",
    "link",
]

# Extensions that are not useful as dataset "formats" in the audit.
_SKIP_EXTENSIONS = frozenset({"", "gitkeep", "ds_store"})


@dataclass
class S3ScanResult:
    size_bytes: int | None = None
    formats: set[str] = field(default_factory=set)


def _ts_to_datetime(ts: int | None) -> str:
    if ts is None:
        return ""
    return datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def _item_to_dict(item: Any) -> dict[str, Any]:
    if isinstance(item, dict):
        return item
    if hasattr(item, "model_dump"):
        return item.model_dump()
    if hasattr(item, "dict"):
        return item.dict()
    return dict(item)


def _nested_get(d: dict[str, Any], *keys: str, default: Any = "") -> Any:
    cur: Any = d
    for k in keys:
        if not isinstance(cur, dict):
            return default
        cur = cur.get(k)
    return cur if cur is not None else default


def _parse_s3_uri(s3_uri: str) -> tuple[str, str]:
    parsed = urlparse(s3_uri)
    bucket = parsed.netloc
    prefix = parsed.path.lstrip("/")
    if prefix and not prefix.endswith("/"):
        prefix += "/"
    return bucket, prefix


def _boto3_client_from_creds(creds_response: Any) -> Any:
    c = creds_response.credentials
    return boto3.client(
        "s3",
        aws_access_key_id=c.aws_access_key_id,
        aws_secret_access_key=c.aws_secret_access_key,
        aws_session_token=c.aws_session_token,
    )


def _extension_from_s3_key(key: str, prefix: str) -> str | None:
    if key.endswith("/"):
        return None
    if prefix and key.startswith(prefix):
        rel = key[len(prefix) :]
    else:
        rel = key
    if not rel or rel.endswith("/"):
        return None
    suffix = PurePosixPath(rel).suffix.lower().lstrip(".")
    if not suffix or suffix in _SKIP_EXTENSIONS:
        return None
    return suffix


def _scan_s3_prefix_sync(s3_client: Any, bucket: str, prefix: str) -> S3ScanResult:
    """List objects under prefix; sum sizes and collect unique file extensions."""
    total = 0
    formats: set[str] = set()
    paginator = s3_client.get_paginator("list_objects_v2")
    for page in paginator.paginate(Bucket=bucket, Prefix=prefix):
        for obj in page.get("Contents", []):
            key = obj.get("Key", "")
            total += int(obj.get("Size", 0))
            ext = _extension_from_s3_key(key, prefix)
            if ext:
                formats.add(ext)
    return S3ScanResult(size_bytes=total, formats=formats)


def _metadata_formats(di: dict[str, Any]) -> str:
    formats = di.get("formats") or []
    if isinstance(formats, list):
        return ", ".join(str(f) for f in formats if f)
    return str(formats) if formats else ""


def _formats_display(s3_formats: set[str], di: dict[str, Any]) -> str:
    if s3_formats:
        return ", ".join(sorted(s3_formats))
    return _metadata_formats(di)


async def _resolve_display_name(client: ProvenaClient, cache: dict[str, str], item_id: str) -> str:
    if not item_id:
        return ""
    if item_id in cache:
        return cache[item_id]
    try:
        result = await client.registry.general_fetch_item(id=item_id)
        if result.status.success and result.item:
            data = _item_to_dict(result.item)
            name = data.get("display_name") or data.get("name") or item_id
            cache[item_id] = str(name)
            return cache[item_id]
    except Exception:
        pass
    cache[item_id] = item_id
    return item_id


async def _build_person_email_cache(client: ProvenaClient) -> dict[str, str]:
    """Map owner email/username -> person display name."""
    email_cache: dict[str, str] = {}
    request = GeneralListRequest(
        filter_by=FilterOptions(
            record_type=QueryRecordTypes.COMPLETE_ONLY,
            item_subtype=ItemSubType.PERSON,
            release_reviewer=None,
            release_status=None,
        ),
        sort_by=SortOptions(sort_type=SortType.DISPLAY_NAME, ascending=True, begins_with=None),
        pagination_key=None,
        page_size=100,
    )
    while True:
        result = await client.registry.person.list_items(list_items_payload=request)
        if result.items:
            for person in result.items:
                data = _item_to_dict(person)
                email = (data.get("email") or "").strip().lower()
                display = data.get("display_name") or ""
                if email and display:
                    email_cache[email] = str(display)
        if not result.pagination_key:
            break
        request.pagination_key = result.pagination_key
    return email_cache


def _record_author(data: dict[str, Any], person_by_email: dict[str, str]) -> str:
    """Record author = registry owner_username, resolved to Person display name when linked."""
    username = (data.get("owner_username") or "").strip()
    if not username:
        history = data.get("history") or []
        if history:
            username = (history[0].get("username") or "").strip()
    if not username:
        return ""
    return person_by_email.get(username.lower(), username)


async def _dataset_s3_scan(
    client: ProvenaClient, dataset_id: str, max_retries: int = 3
) -> S3ScanResult:
    """Fetch S3 credentials and scan prefix for total size and file extensions."""
    last_error: Exception | None = None
    for attempt in range(max_retries):
        try:
            fetch = await client.datastore.fetch_dataset(id=dataset_id)
            if not fetch.item or not fetch.item.s3:
                return S3ScanResult()
            s3_uri = fetch.item.s3.s3_uri
            bucket, prefix = _parse_s3_uri(s3_uri)

            creds = await client.datastore.generate_read_access_credentials(
                credentials=CredentialsRequest(
                    dataset_id=dataset_id, console_session_required=False
                )
            )
            s3_client = _boto3_client_from_creds(creds)
            return await asyncio.to_thread(_scan_s3_prefix_sync, s3_client, bucket, prefix)
        except Exception as exc:
            last_error = exc
            if attempt < max_retries - 1:
                await asyncio.sleep(2**attempt)
    if last_error:
        print(f"    warning: S3 scan failed for {dataset_id}: {last_error}", flush=True)
    return S3ScanResult()


async def _scan_datasets_s3(
    client: ProvenaClient,
    dataset_ids: list[str],
    *,
    label: str,
    concurrency: int,
    max_retries: int,
) -> dict[str, S3ScanResult]:
    results: dict[str, S3ScanResult] = {}
    sem = asyncio.Semaphore(concurrency)
    total = len(dataset_ids)

    async def _task(ds_id: str, index: int) -> tuple[str, S3ScanResult]:
        async with sem:
            print(f"  [{index}/{total}] {ds_id} ({label})", flush=True)
            return ds_id, await _dataset_s3_scan(client, ds_id, max_retries=max_retries)

    gathered = await asyncio.gather(*[_task(ds_id, i) for i, ds_id in enumerate(dataset_ids, 1)])
    results.update(gathered)
    return results


def _excel_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _format_bytes(num: int | None) -> str:
    if num is None:
        return ""
    if num < 1024:
        return f"{num} B"
    if num < 1024**2:
        return f"{num / 1024:.2f} KB"
    if num < 1024**3:
        return f"{num / 1024**2:.2f} MB"
    return f"{num / 1024**3:.2f} GB"


async def _list_all_datasets(client: ProvenaClient) -> list[dict[str, Any]]:
    """List all datasets via the datastore API, sorted by modified time descending."""
    sort = SortOptions(sort_type=SortType.UPDATED_TIME, ascending=False, begins_with=None)
    request = NoFilterSubtypeListRequest(sort_by=sort, pagination_key=None, page_size=100)
    datasets: list[dict[str, Any]] = []
    while True:
        result = await client.datastore.list_datasets(list_dataset_request=request)
        if result.items:
            for item in result.items:
                datasets.append(_item_to_dict(item))
        if not result.pagination_key:
            break
        request.pagination_key = result.pagination_key
    datasets.sort(key=lambda d: d.get("updated_timestamp") or 0, reverse=True)
    return datasets


def _merge_scan(prev: S3ScanResult, new: S3ScanResult) -> S3ScanResult:
    return S3ScanResult(
        size_bytes=new.size_bytes if new.size_bytes is not None else prev.size_bytes,
        formats=prev.formats | new.formats,
    )


def _ids_missing_size(
    prepared: list[tuple[dict[str, Any], str, str, bool]],
    scan_by_id: dict[str, S3ScanResult],
) -> list[str]:
    """Dataset IDs that still have no S3 size after a scan pass."""
    missing: list[str] = []
    for data, _, _, _ in prepared:
        ds_id = data.get("id", "")
        if not ds_id:
            continue
        scan = scan_by_id.get(ds_id, S3ScanResult())
        if scan.size_bytes is None:
            missing.append(ds_id)
    return missing


def _row_from_dataset(
    data: dict[str, Any],
    author_name: str,
    custodian_name: str,
    scan: S3ScanResult | None,
    reposited: bool,
) -> list[Any]:
    ds_id = data.get("id", "")
    cf = data.get("collection_format") or {}
    di = cf.get("dataset_info") or {}

    s3_formats = scan.formats if scan else set()
    format_str = _formats_display(s3_formats, di)

    created = _ts_to_datetime(data.get("created_timestamp"))
    modified = _ts_to_datetime(data.get("updated_timestamp"))
    version_info = data.get("versioning_info") or {}
    version = version_info.get("version", "")

    size_bytes = scan.size_bytes if scan else None
    if reposited or size_bytes is not None:
        size_display = _format_bytes(size_bytes)
    else:
        size_display = ""

    return [
        ds_id,
        data.get("display_name") or di.get("name") or "",
        di.get("description") or "",
        format_str,
        created,
        modified,
        str(di.get("license") or ""),
        size_display,
        author_name,
        custodian_name,
        str(version) if version != "" else "",
        f"{HANDLE_URL_PREFIX}{ds_id}" if ds_id else "",
    ]


async def run(output: Path, skip_size: bool) -> None:
    load_dotenv(_project_root / ".env", override=False)
    cfg = pr.get_provena_config()
    token = (cfg.get("token") or "").strip()
    if not token:
        raise SystemExit(
            "No offline token found. Set PROVENA_OFFLINE_TOKEN or add a token to provena_tokens.json."
        )

    api_overrides = pr.build_api_overrides_for_config(cfg)
    config = Config(
        domain=str(cfg.get("domain") or "dev.rrap-is.com"),
        realm_name=str(cfg.get("realm") or "rrap"),
        api_overrides=api_overrides,
    )
    auth = OfflineFlow(
        config=config,
        client_id=os.getenv("MCP_OFFLINE_CLIENT_ID", "automated-access"),
        offline_token=token,
        log_level=Log.ERROR,
    )
    client = ProvenaClient(auth=auth, config=config)

    instance = cfg.get("instance") or config.domain
    print(f"Listing datasets on instance: {instance}")
    datasets = await _list_all_datasets(client)
    print(f"Found {len(datasets)} datasets (sorted by date modified, newest first)")

    print("Loading person registry for record author resolution...")
    person_by_email = await _build_person_email_cache(client)

    name_cache: dict[str, str] = {}
    prepared: list[tuple[dict[str, Any], str, str, bool]] = []

    for data in datasets:
        cf = data.get("collection_format") or {}
        assoc = cf.get("associations") or {}
        author_name = _record_author(data, person_by_email)
        custodian_id = assoc.get("data_custodian_id") or ""
        custodian_name = await _resolve_display_name(client, name_cache, custodian_id)
        reposited = bool(_nested_get(cf.get("dataset_info") or {}, "access_info", "reposited", default=False))
        prepared.append((data, author_name, custodian_name, reposited))

    scan_by_id: dict[str, S3ScanResult] = {}
    if not skip_size:
        reposited_ids = [d.get("id", "") for d, _, _, r in prepared if r and d.get("id")]
        print(f"Scanning AWS S3 ({len(reposited_ids)} reposited datasets: size + formats)...")
        scan_by_id = await _scan_datasets_s3(
            client, reposited_ids, label="initial", concurrency=12, max_retries=3
        )

        for pass_label, concurrency, max_retries in (
            ("retry-1", 4, 5),
            ("retry-2", 2, 8),
            ("retry-3", 1, 10),
        ):
            missing_size = _ids_missing_size(prepared, scan_by_id)
            if not missing_size:
                break
            print(
                f"Retrying {len(missing_size)} datasets with missing size "
                f"({pass_label}, concurrency={concurrency})..."
            )
            retry_results = await _scan_datasets_s3(
                client,
                missing_size,
                label=pass_label,
                concurrency=concurrency,
                max_retries=max_retries,
            )
            for ds_id, result in retry_results.items():
                prev = scan_by_id.get(ds_id, S3ScanResult())
                scan_by_id[ds_id] = _merge_scan(prev, result)

        still_missing = len(_ids_missing_size(prepared, scan_by_id))
        total = len(prepared)
        print(f"Size populated for {total - still_missing}/{total} datasets")

    prepared.sort(key=lambda item: item[0].get("updated_timestamp") or 0, reverse=True)

    rows: list[list[Any]] = []
    for data, author_name, custodian_name, reposited in prepared:
        ds_id = data.get("id", "")
        scan = scan_by_id.get(ds_id) if not skip_size else None
        rows.append(_row_from_dataset(data, author_name, custodian_name, scan, reposited))

    wb = Workbook()
    ws = wb.active
    ws.title = "Datasets"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_excel_value(v) for v in row])

    for col_idx, header in enumerate(COLUMNS, 1):
        letter = get_column_letter(col_idx)
        max_len = len(header)
        for row in ws.iter_rows(
            min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx, values_only=True
        ):
            val = row[0]
            if val is not None:
                max_len = max(max_len, min(len(str(val)), 80))
        ws.column_dimensions[letter].width = max_len + 2

    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(output)
        print(f"Wrote {len(rows)} rows to {output}")
    except PermissionError:
        alt = output.with_name(f"{output.stem}_updated{output.suffix}")
        wb.save(alt)
        print(f"Could not overwrite {output} (file may be open in Excel). Wrote {len(rows)} rows to {alt}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export Provena datasets to Excel audit spreadsheet.")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=_project_root / "mds_dataset_audit.xlsx",
        help="Output .xlsx path (default: mds_dataset_audit.xlsx in repo root)",
    )
    parser.add_argument(
        "--skip-size",
        action="store_true",
        help="Skip AWS S3 scan (faster; size and S3-derived formats left blank for reposited data)",
    )
    args = parser.parse_args()
    asyncio.run(run(args.output.resolve(), skip_size=args.skip_size))


if __name__ == "__main__":
    main()
